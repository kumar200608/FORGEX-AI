import os
import json
import io
from pathlib import Path
from PIL import Image
import pymupdf

def bbox_overlap(box1, box2, threshold=0.4):
    """
    Returns True if box1 is significantly inside box2.
    box format: [x0, y0, x1, y1]
    """
    x0 = max(box1[0], box2[0])
    y0 = max(box1[1], box2[1])
    x1 = min(box1[2], box2[2])
    y1 = min(box1[3], box2[3])

    if x1 <= x0 or y1 <= y0:
        return False

    intersection = (x1 - x0) * (y1 - y0)
    area1 = (box1[2] - box1[0]) * (box1[3] - box1[1])
    if area1 <= 0:
        return False

    return (intersection / area1) >= threshold

def clean_table_markdown(cells):
    """
    Converts 2D cell grid into clean, standard Markdown table.
    """
    if not cells:
        return ""
    
    cleaned = []
    for row in cells:
        cleaned.append([(c or "").strip().replace("\n", " ") for c in row])
        
    # Filter out empty rows
    cleaned = [r for r in cleaned if any(cell for cell in r)]
    if not cleaned:
        return ""

    num_cols = max(len(r) for r in cleaned)
    if num_cols == 0:
        return ""

    # Pad all rows to num_cols
    normalized = []
    for r in cleaned:
        padded = r + [""] * (num_cols - len(r))
        normalized.append([c.replace("|", "\\|") for c in padded])

    headers = normalized[0]
    data_rows = normalized[1:] if len(normalized) > 1 else []

    header_line = "| " + " | ".join(headers) + " |"
    separator_line = "| " + " | ".join(["---"] * num_cols) + " |"
    data_lines = ["| " + " | ".join(r) + " |" for r in data_rows]

    return "\n".join([header_line, separator_line] + data_lines)


class Pdf2Json:
    def __init__(self, name):
        """
        name can be a UUID, filename with or without extension.
        """
        self.name = name
        self.base_name = os.path.splitext(os.path.basename(name))[0]
        
        # Base directories
        self.downloads_dir = Path("downloads")
        self.output_dir = Path("langbase_json")
        self.images_dir = self.output_dir / "ExtractedImages" / self.base_name

        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.images_dir.mkdir(parents=True, exist_ok=True)

        # Locate PDF file
        backend_downloads = Path(__file__).resolve().parent / "downloads"
        candidates = [
            self.downloads_dir / f"{self.base_name}.pdf",
            self.downloads_dir / f"{name}.pdf",
            backend_downloads / f"{self.base_name}.pdf",
            backend_downloads / f"{name}.pdf",
            Path(f"{name}.pdf"),
            Path(name),
            self.downloads_dir / name,
        ]
        self.pdf_path = None
        for c in candidates:
            if c.exists() and not c.is_dir() and c.suffix.lower() == ".pdf":
                self.pdf_path = c
                break

        if not self.pdf_path:
            # Fallback search in downloads directory
            for p in self.downloads_dir.glob(f"{self.base_name}*"):
                if p.suffix.lower() == ".pdf":
                    self.pdf_path = p
                    break

        # Check for PPTX
        self.pptx_path = None
        if not self.pdf_path:
            for p in self.downloads_dir.glob(f"{self.base_name}*"):
                if p.suffix.lower() in [".pptx", ".ppt"]:
                    self.pptx_path = p
                    break

    def _consolidate_text_blocks(self, blocks, page_idx):
        """
        Consolidates fragmented lines of the same paragraph into cohesive text blocks
        with precise union bounding boxes.
        """
        if not blocks:
            return []

        sorted_blocks = sorted(blocks, key=lambda b: (b["bbox"][1], b["bbox"][0]))
        merged = []
        curr = None

        for b in sorted_blocks:
            if curr is None:
                curr = dict(b)
                continue

            c_box = curr["bbox"]
            b_box = b["bbox"]

            v_gap = b_box[1] - c_box[3]
            h_overlap = max(0, min(c_box[2], b_box[2]) - max(c_box[0], b_box[0]))
            same_left = abs(c_box[0] - b_box[0]) < 35

            c_text = curr["text"].strip()
            b_text = b["text"].strip()

            is_curr_heading = len(c_text) < 30 and (c_text.isupper() or c_text.startswith("Ref") or "CERTIFICATE" in c_text)
            is_b_heading = len(b_text) < 30 and (b_text.isupper() or b_text.startswith("Ref") or "CERTIFICATE" in b_text or "Principal" in b_text or "Signatory" in b_text)

            can_merge = (
                -4 <= v_gap <= 16 and
                (h_overlap > 30 or same_left) and
                not is_curr_heading and
                not is_b_heading
            )

            if can_merge:
                curr["bbox"] = [
                    round(min(c_box[0], b_box[0]), 1),
                    round(min(c_box[1], b_box[1]), 1),
                    round(max(c_box[2], b_box[2]), 1),
                    round(max(c_box[3], b_box[3]), 1),
                ]
                curr["text"] = c_text + " " + b_text
            else:
                merged.append(curr)
                curr = dict(b)

        if curr:
            merged.append(curr)

        for idx, blk in enumerate(merged, start=1):
            blk["block_id"] = f"p{page_idx}_b{idx}"

        return merged

    def extract(self):
        if self.pdf_path and self.pdf_path.exists():
            return self.extract_pdf()
        elif self.pptx_path and self.pptx_path.exists():
            return self.extract_pptx()
        else:
            print(f"[WARN] No PDF or PPTX found for {self.name} in downloads/")
            return {}

    def extract_pdf(self):
        print(f"[INFO] Extracting rich mixed-format content from: {self.pdf_path}")
        doc = pymupdf.open(str(self.pdf_path))

        legacy_format = {}
        image_counter = 1

        for page_idx, page in enumerate(doc, start=1):
            page_key = f"page_{page_idx}"
            page_rect = page.rect
            width, height = float(page_rect.width), float(page_rect.height)

            tables_meta = []
            table_bboxes = []

            # ------------------------------------------------------------------
            # 1. TABLE DETECTION & EXTRACTION (PyMuPDF find_tables)
            # ------------------------------------------------------------------
            try:
                tabs = page.find_tables()
                for t_idx, tab in enumerate(tabs, start=1):
                    t_bbox = [round(float(c), 1) for c in tab.bbox]
                    table_bboxes.append(t_bbox)
                    cells = tab.extract()
                    md_table = clean_table_markdown(cells)
                    if md_table:
                        tables_meta.append({
                            "table_id": f"p{page_idx}_t{t_idx}",
                            "bbox": t_bbox,
                            "markdown": md_table,
                            "num_rows": len(cells) if cells else 0,
                            "num_cols": len(cells[0]) if cells and cells[0] else 0,
                            "headers": cells[0] if cells else []
                        })
            except Exception as e:
                print(f"[WARN] Table detection error on page {page_idx}: {e}")

            # ------------------------------------------------------------------
            # 2. TEXT BLOCKS WITH SPATIAL BOUNDING BOXES
            # ------------------------------------------------------------------
            raw_text_blocks = []
            raw_blocks = page.get_text("blocks")
            for b_idx, block in enumerate(raw_blocks):
                # block tuple: (x0, y0, x1, y1, text, block_no, block_type)
                if len(block) >= 5:
                    x0, y0, x1, y1, b_text = block[:5]
                    b_type = block[6] if len(block) > 6 else 0
                    if b_type == 0:  # text
                        b_text_clean = b_text.strip()
                        if len(b_text_clean) > 3:
                            b_bbox = [
                                round(float(x0), 1),
                                round(float(y0), 1),
                                round(float(x1), 1),
                                round(float(y1), 1)
                            ]
                            # Check if this block falls inside a detected table
                            is_in_table = any(bbox_overlap(b_bbox, tb) for tb in table_bboxes)
                            if not is_in_table:
                                raw_text_blocks.append({
                                    "block_id": f"p{page_idx}_b{b_idx}",
                                    "bbox": b_bbox,
                                    "text": b_text_clean
                                })

            # Consolidate consecutive paragraph lines into cohesive semantic blocks
            text_blocks = self._consolidate_text_blocks(raw_text_blocks, page_idx)

            # ------------------------------------------------------------------
            # 3. IMAGES, CHARTS & VISUAL CROPS
            # ------------------------------------------------------------------
            images_meta = []
            try:
                image_list = page.get_images(full=True)
                for img_info in image_list:
                    xref = img_info[0]
                    rects = page.get_image_rects(xref)
                    if rects:
                        img_rect = rects[0]
                        if img_rect.width >= 40 and img_rect.height >= 40:
                            img_bbox = [round(float(c), 1) for c in img_rect]
                            img_filename = f"{self.base_name}_p{page_idx}_img{image_counter}.png"
                            img_save_path = self.images_dir / img_filename

                            pix = page.get_pixmap(clip=img_rect, dpi=150)
                            pix.save(str(img_save_path))

                            is_chart = (img_rect.width > 120 and img_rect.height > 80)
                            images_meta.append({
                                "image_id": f"p{page_idx}_i{image_counter}",
                                "filename": img_filename,
                                "path": str(img_save_path),
                                "bbox": img_bbox,
                                "is_chart": is_chart,
                                "width": round(float(img_rect.width), 1),
                                "height": round(float(img_rect.height), 1)
                            })
                            image_counter += 1
            except Exception as e:
                print(f"[WARN] Image extraction error on page {page_idx}: {e}")

            # ------------------------------------------------------------------
            # 4. ASSEMBLE PAGE RECORD
            # ------------------------------------------------------------------
            page_data = {
                "page_num": page_idx,
                "width": round(width, 1),
                "height": round(height, 1),
                "tables": tables_meta,
                "text_blocks": text_blocks,
                "images": images_meta
            }

            combined_page_text = []
            for tb in tables_meta:
                combined_page_text.append(f"[Table {tb['table_id']}]\n{tb['markdown']}")
            for blk in text_blocks:
                combined_page_text.append(blk["text"])

            legacy_format[page_key] = {
                "text": "\n\n".join(combined_page_text),
                "images": [im["filename"] for im in images_meta],
                "page_num": page_idx,
                "tables": tables_meta,
                "text_blocks": text_blocks,
                "raw_page_data": page_data
            }

        doc.close()

        # Save to langbase_json/{name}.json
        output_json_path = self.output_dir / f"{self.base_name}.json"
        with open(output_json_path, "w", encoding="utf-8") as f:
            json.dump(legacy_format, f, indent=2, ensure_ascii=False)

        print(f"[SUCCESS] Extracted {len(legacy_format)} pages with tables & bounding boxes to: {output_json_path}")
        return legacy_format

    def extract_pptx(self):
        """Native extraction for PowerPoint files without proprietary Aspose."""
        print(f"[INFO] Extracting PowerPoint content using python-pptx: {self.pptx_path}")
        try:
            from pptx import Presentation
            prs = Presentation(str(self.pptx_path))
            slide_w = prs.slide_width.pt if hasattr(prs, "slide_width") else 720.0
            slide_h = prs.slide_height.pt if hasattr(prs, "slide_height") else 540.0

            legacy_format = {}

            for s_idx, slide in enumerate(prs.slides, start=1):
                page_key = f"page_{s_idx}"
                tables_meta = []
                text_blocks = []

                for shape_idx, shape in enumerate(slide.shapes):
                    # Bounding box in points (72 DPI)
                    left = float(shape.left.pt) if hasattr(shape, "left") else 0.0
                    top = float(shape.top.pt) if hasattr(shape, "top") else 0.0
                    width = float(shape.width.pt) if hasattr(shape, "width") else 0.0
                    height = float(shape.height.pt) if hasattr(shape, "height") else 0.0
                    bbox = [round(left, 1), round(top, 1), round(left + width, 1), round(top + height, 1)]

                    if shape.has_table:
                        cells = []
                        for row in shape.table.rows:
                            cells.append([cell.text.strip() for cell in row.cells])
                        md = clean_table_markdown(cells)
                        if md:
                            tables_meta.append({
                                "table_id": f"p{s_idx}_t{shape_idx}",
                                "bbox": bbox,
                                "markdown": md,
                                "headers": cells[0] if cells else []
                            })
                    elif shape.has_text_frame:
                        text = shape.text.strip()
                        if len(text) > 3:
                            text_blocks.append({
                                "block_id": f"p{s_idx}_b{shape_idx}",
                                "bbox": bbox,
                                "text": text
                            })

                combined = []
                for tb in tables_meta:
                    combined.append(f"[Table {tb['table_id']}]\n{tb['markdown']}")
                for blk in text_blocks:
                    combined.append(blk["text"])

                legacy_format[page_key] = {
                    "text": "\n\n".join(combined),
                    "images": [],
                    "page_num": s_idx,
                    "tables": tables_meta,
                    "text_blocks": text_blocks,
                    "raw_page_data": {
                        "page_num": s_idx,
                        "width": slide_w,
                        "height": slide_h,
                        "tables": tables_meta,
                        "text_blocks": text_blocks,
                        "images": []
                    }
                }

            output_json_path = self.output_dir / f"{self.base_name}.json"
            with open(output_json_path, "w", encoding="utf-8") as f:
                json.dump(legacy_format, f, indent=2, ensure_ascii=False)

            print(f"[SUCCESS] Extracted {len(legacy_format)} slides to: {output_json_path}")
            return legacy_format
        except Exception as e:
            print(f"[ERROR] Error extracting PPTX: {e}")
            return {}

if __name__ == "__main__":
    import sys
    name = sys.argv[1] if len(sys.argv) > 1 else "sample"
    obj = Pdf2Json(name)
    res = obj.extract()
    print("Done. Pages extracted:", len(res))
