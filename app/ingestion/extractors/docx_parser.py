"""Microsoft Word (DOCX) document extractor for TraceGuard AI."""

import io
from typing import List
import docx

from app.ingestion.extractors.base import BaseExtractor, ExtractionResult, ExtractionStatus


class DocxExtractor(BaseExtractor):
    """Safely extracts text, tables, and structure from DOCX without macro execution."""

    @property
    def supported_extensions(self) -> List[str]:
        return [".docx"]

    @property
    def format_name(self) -> str:
        return "Microsoft Word (DOCX)"

    def extract(self, filename: str, content: bytes) -> ExtractionResult:
        # DOCX files are zip archives starting with PK (0x50 0x4B 0x03 0x04)
        if not content.startswith(b"PK"):
            return ExtractionResult(
                status=ExtractionStatus.FAILED,
                error_message="Invalid DOCX signature: file does not match zip container header.",
            )

        try:
            doc = docx.Document(io.BytesIO(content))
            sections_text = []

            # 1. Paragraphs
            para_texts = []
            for p in doc.paragraphs:
                t = p.text.strip()
                if t:
                    para_texts.append(t)
            if para_texts:
                sections_text.append("\n".join(para_texts))

            # 2. Tables
            table_count = len(doc.tables)
            table_texts = []
            for t_idx, table in enumerate(doc.tables):
                t_lines = [f"[Table {t_idx + 1}]"]
                for row in table.rows:
                    row_vals = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    # Filter consecutive duplicates created by cell merges
                    deduped_row = []
                    for val in row_vals:
                        if not deduped_row or val != deduped_row[-1]:
                            deduped_row.append(val)
                    if any(deduped_row):
                        t_lines.append(" | ".join(deduped_row))
                if len(t_lines) > 1:
                    table_texts.append("\n".join(t_lines))

            if table_texts:
                sections_text.append("\n\n".join(table_texts))

            full_text = "\n\n".join(sections_text).strip()

            metadata = {
                "paragraph_count": len(para_texts),
                "table_count": table_count,
                "has_macros": "vbaProject.bin" in str(content[:4096]),
            }

            warnings = []
            if metadata["has_macros"]:
                warnings.append("Document contains VBA macro streams (macros were suppressed).")

            if not full_text:
                return ExtractionResult(
                    status=ExtractionStatus.PARTIAL,
                    raw_text="",
                    metadata=metadata,
                    warnings=warnings + ["Document contains no extractable text paragraphs or tables."],
                )

            return ExtractionResult(
                status=ExtractionStatus.SUCCESS,
                raw_text=full_text,
                metadata=metadata,
                warnings=warnings,
            )
        except Exception as e:
            return ExtractionResult(
                status=ExtractionStatus.FAILED,
                error_message=f"DOCX parsing error: {str(e)}",
            )
